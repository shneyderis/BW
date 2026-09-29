#!/usr/bin/env python3
"""
Конвертація вивантаження 1С (beykush_data.xlsx) → data/1c_*.csv + data/manifest.json

Використання (з кореня репозиторію):
    python3 tools/convert_1c.py                # xlsx → csv + manifest
    python3 tools/convert_1c.py --lots-only    # лише оновити manifest для data/lots.csv (після заміни файлу)
    python3 tools/convert_1c.py --manifest     # лише перерахувати manifest.json по наявних файлах

Потрібно: pip install openpyxl (тільки для конвертації xlsx).
manifest.json читає дашборд (js/health.js) — щоб показувати дату вивантаження файлів і задачі "оновити".
"""
import csv, json, os, re, sys, datetime

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
XLSX = os.path.join(ROOT, "beykush_data.xlsx")
DATA = os.path.join(ROOT, "data")
MANIFEST = os.path.join(DATA, "manifest.json")

SHEETS = {"Продажи": "sales", "Контрагенты": "partners", "Номенклатура": "products",
          "ОСВ": "osv", "Банк": "bank", "Склады": "warehouses", "Договоры": "contracts",
          "Организации": "orgs", "Сотрудники": "staff", "ОС": "assets"}
DATE_COLS = {"sales": "Дата", "bank": "Дата", "contracts": "Дата"}


def today():
    return datetime.date.today().isoformat()


def convert_xlsx():
    try:
        import openpyxl
    except ImportError:
        sys.exit("Потрібен openpyxl: pip install openpyxl")
    if not os.path.exists(XLSX):
        sys.exit(f"Немає файла {XLSX}")
    wb = openpyxl.load_workbook(XLSX, read_only=True, data_only=True)
    os.makedirs(DATA, exist_ok=True)
    for name, eng in SHEETS.items():
        if name not in wb.sheetnames:
            print(f"  ! аркуш «{name}» відсутній — пропускаю")
            continue
        ws = wb[name]
        rows = list(ws.iter_rows(values_only=True))
        hdr_idx = 3 if name == "ОСВ" else 2
        headers = [str(c).strip() if c else f"col{i}" for i, c in enumerate(rows[hdr_idx])]
        fname = os.path.join(DATA, f"1c_{eng}.csv")
        n = 0
        with open(fname, "w", newline="", encoding="utf-8") as f:
            w = csv.writer(f)
            w.writerow(headers)
            for r in rows[hdr_idx + 1:]:
                if any(c for c in r):
                    w.writerow([str(c) if c is not None else "" for c in r])
                    n += 1
        print(f"  {name} → data/1c_{eng}.csv ({n} рядків)")


def max_date(path, col):
    """Максимальна дата у колонці (dd.mm.yyyy або yyyy-mm-dd) → 'YYYY-MM-DD'."""
    best = ""
    with open(path, encoding="utf-8") as f:
        r = csv.DictReader(f)
        if col not in (r.fieldnames or []):
            return ""
        for row in r:
            v = (row.get(col) or "").strip()
            m = re.match(r"(\d{2})\.(\d{2})\.(\d{4})", v)
            iso = f"{m.group(3)}-{m.group(2)}-{m.group(1)}" if m else (v[:10] if re.match(r"\d{4}-\d{2}-\d{2}", v) else "")
            if iso and iso <= today() and iso > best:   # дати з майбутнього ігноруємо
                best = iso
    return best


def build_manifest(touched=None, lots_only=False):
    """touched — файли, для яких проставити exported=сьогодні; інші беруть попередню дату з manifest."""
    old = {}
    if os.path.exists(MANIFEST):
        try:
            old = json.load(open(MANIFEST, encoding="utf-8")).get("files", {})
        except Exception:
            old = {}
    files = {}
    names = ["data/lots.csv"] if lots_only else [f"data/1c_{e}.csv" for e in SHEETS.values()] + ["data/lots.csv"]
    for rel in sorted(set(names) | set(old.keys())):
        path = os.path.join(ROOT, rel)
        if not os.path.exists(path):
            continue
        with open(path, encoding="utf-8", newline="") as f:
            rows = max(0, sum(1 for _ in csv.reader(f)) - 1)
        eng = re.sub(r"^data/1c_|\.csv$", "", rel)
        entry = {"exported": (today() if (touched is None or rel in touched) else old.get(rel, {}).get("exported", today())),
                 "rows": rows}
        if eng in DATE_COLS:
            md = max_date(path, DATE_COLS[eng])
            if md:
                entry["max_date"] = md
        files[rel] = entry
    out = {"generated": today(), "files": files}
    json.dump(out, open(MANIFEST, "w", encoding="utf-8"), ensure_ascii=False, indent=1)
    print(f"  manifest.json: {len(files)} файлів, generated {out['generated']}")


if __name__ == "__main__":
    args = sys.argv[1:]
    if "--manifest" in args:
        build_manifest(touched=set())            # дати не чіпаємо
    elif "--lots-only" in args:
        build_manifest(touched={"data/lots.csv"})
    else:
        print("Конвертація beykush_data.xlsx …")
        convert_xlsx()
        build_manifest(touched={f"data/1c_{e}.csv" for e in SHEETS.values()})
    print("Готово. Далі: git add data && git commit -m 'Update 1C data' && git push")
