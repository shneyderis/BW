#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Витягує листи виписок (PRIVAT, PRIVAT21-24, VOSTOK) з xlsx-експорту книги
BW_accounts у data/bank_raw/sheet_*.csv для build_postings_daily.py.
Запуск: /usr/bin/python3.13 accounting/extract_book_sheets.py <шлях до bw_accounts.xlsx>
(книгу BW_accounts експортувати як xlsx: Drive → download as xlsx)
"""
import sys, csv, os
import openpyxl

OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "data", "bank_raw")
SHEETS = {"PRIVAT": "sheet_privat_current.csv",
          "PRIVAT21-24": "sheet_privat_2021-24.csv",
          "VOSTOK": "sheet_vostok.csv"}

def cell(v):
    if v is None: return ""
    if hasattr(v, "strftime"): return v.strftime("%Y-%m-%d %H:%M:%S")
    return str(v)

def main(path):
    wb = openpyxl.load_workbook(path, read_only=True)
    for sh, fname in SHEETS.items():
        if sh not in wb.sheetnames:
            print(f"  {sh}: немає в книзі"); continue
        ws = wb[sh]
        rows = [[cell(c) for c in r] for r in ws.iter_rows(values_only=True)]
        with open(os.path.join(OUT, fname), "w", encoding="utf-8", newline="") as f:
            csv.writer(f).writerows(rows)
        print(f"  {sh} → {fname}: {len(rows)} рядків")

if __name__ == "__main__":
    main(sys.argv[1])
